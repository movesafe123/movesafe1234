import sys
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Set inner padding for a cell in dxa (20 dxa = 1 pt)."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_valign(cell, valign='center'):
    """Set vertical alignment of cell."""
    tcPr = cell._element.get_or_add_tcPr()
    # Remove existing vAlign if any
    for existing in tcPr.findall(qn('w:vAlign')):
        tcPr.remove(existing)
    vAlign = OxmlElement('w:vAlign')
    vAlign.set(qn('w:val'), valign)
    tcPr.append(vAlign)

def set_cell_width(cell, width_dxa):
    """Set exact width of cell in dxa."""
    tcPr = cell._element.get_or_add_tcPr()
    for existing in tcPr.findall(qn('w:tcW')):
        tcPr.remove(existing)
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(width_dxa))
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)

def set_table_borders(table):
    """Set clean, professional solid borders for the entire table."""
    tblPr = table._element.tblPr
    for existing in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(existing)
    
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_table_layout_and_center(table, col_widths_dxa):
    """Set table alignment to center, fixed layout, cell margins and column grid."""
    tblPr = table._element.tblPr
    
    # Remove existing jc, tblW, tblCellMar, tblLayout
    for tag in ['w:jc', 'w:tblW', 'w:tblCellMar', 'w:tblLayout']:
        for existing in tblPr.findall(qn(tag)):
            tblPr.remove(existing)
            
    total_w = sum(col_widths_dxa)
    
    # Table width
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(total_w))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)
    
    # Table align center
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    tblPr.append(jc)
    
    # Fixed layout
    tblLayout = OxmlElement('w:tblLayout')
    tblLayout.set(qn('w:type'), 'fixed')
    tblPr.append(tblLayout)
    
    # Default cell margins
    tblCellMar = parse_xml(
        f'<w:tblCellMar {nsdecls("w")}>\n'
        f'  <w:top w:w="100" w:type="dxa"/>\n'
        f'  <w:left w:w="160" w:type="dxa"/>\n'
        f'  <w:bottom w:w="100" w:type="dxa"/>\n'
        f'  <w:right w:w="160" w:type="dxa"/>\n'
        f'</w:tblCellMar>'
    )
    tblPr.append(tblCellMar)
    
    # tblGrid
    for existing in table._element.findall(qn('w:tblGrid')):
        table._element.remove(existing)
    
    tblGrid = OxmlElement('w:tblGrid')
    for w in col_widths_dxa:
        gridCol = OxmlElement('w:gridCol')
        gridCol.set(qn('w:w'), str(w))
        tblGrid.append(gridCol)
    # Insert tblGrid right after tblPr
    table._element.insert(table._element.index(tblPr) + 1, tblGrid)

def format_cell_paragraph(p, align='left', bold=False, font_size_pt=12):
    """
    Format paragraph with EXPLICIT alignment (left or center), NEVER justify.
    This eliminates Microsoft Word text stretching / spacing bug completely.
    """
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == 'center' else WD_ALIGN_PARAGRAPH.LEFT
    
    # Explicitly ensure pPr has jc
    pPr = p._element.get_or_add_pPr()
    for existing in pPr.findall(qn('w:jc')):
        pPr.remove(existing)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center' if align == 'center' else 'left')
    pPr.append(jc)
    
    # Set tight, clean spacing
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    
    # Format all runs
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size_pt)
        if bold is not None:
            run.font.bold = bold

def fix_all_tables_in_doc(file_path):
    print(f"Processing {file_path}...")
    doc = docx.Document(file_path)
    
    # Definitions of column widths in dxa (~9640 dxa total ~ 17.0 cm)
    # 1 cm = 567 dxa
    TABLE_CONFIGS = {
        0: {
            "name": "Bảng thông tin thành viên (Table 1)",
            "widths": [650, 2150, 1500, 2600, 1440, 2300], # Total = 10640 dxa
            "col_aligns": ['center', 'left', 'center', 'left', 'center', 'left'],
            "col_bold": [False, False, False, False, False, False],
            "is_team_table": True,
            "padding": {"top": 100, "bottom": 100, "left": 80, "right": 80},
            "font_size": 11
        },
        1: {
            "name": "Bảng 1. Các nhóm dữ liệu của dự án (Table 2)",
            "widths": [2400, 3900, 3340], # Total = 9640 dxa (~17.0cm)
            "col_aligns": ['left', 'left', 'left'],
            "col_bold": [True, False, False]
        },
        2: {
            "name": "Bảng 2. Các bước xử lý dữ liệu (Table 3)",
            "widths": [2800, 6840], # Total = 9640 dxa (~17.0cm)
            "col_aligns": ['left', 'left'],
            "col_bold": [True, False]
        },
        3: {
            "name": "Bảng 3. Công nghệ sử dụng và trạng thái (Table 4)",
            "widths": [2300, 5940, 1400], # Total = 9640 dxa (~17.0cm)
            "col_aligns": ['left', 'left', 'center'],
            "col_bold": [True, False, False]
        },
        4: {
            "name": "Bảng 4. Đánh giá chi tiết (Table 5)",
            "widths": [2000, 7640], # Total = 9640 dxa (~17.0cm)
            "col_aligns": ['left', 'left'],
            "col_bold": [True, False]
        },
        5: {
            "name": "Bảng 5. Ba giai đoạn phát triển dự kiến (Table 6)",
            "widths": [2400, 3900, 3340], # Total = 9640 dxa (~17.0cm)
            "col_aligns": ['left', 'left', 'left'],
            "col_bold": [True, False, False]
        }
    }
    
    for t_idx, table in enumerate(doc.tables):
        cfg = TABLE_CONFIGS.get(t_idx)
        if not cfg:
            continue
        print(f"  Fixing {cfg['name']}...")
        widths = cfg['widths']
        
        # 1. Setup table layout, width, centering and borders
        set_table_borders(table)
        set_table_layout_and_center(table, widths)
        
        # 2. Fix rows
        for r_idx, row in enumerate(table.rows):
            # Row cantSplit and header repeat
            trPr = row._element.get_or_add_trPr()
            if not trPr.findall(qn('w:cantSplit')):
                trPr.append(OxmlElement('w:cantSplit'))
            if r_idx == 0 and not trPr.findall(qn('w:tblHeader')):
                trPr.append(OxmlElement('w:tblHeader'))
                
            for c_idx, cell in enumerate(row.cells):
                col_w = widths[c_idx] if c_idx < len(widths) else 2000
                set_cell_width(cell, col_w)
                set_cell_valign(cell, 'center')
                pad = cfg.get('padding', {'top': 100, 'bottom': 100, 'left': 160, 'right': 160})
                set_cell_margins(cell, top=pad['top'], bottom=pad['bottom'], left=pad['left'], right=pad['right'])
                
                # Special cleanups
                if cfg.get('is_team_table'):
                    # Table 1 header TT fix
                    if r_idx == 0 and c_idx == 0:
                        cell.text = 'TT'
                    pass
                
                # Table 6 giai đoạn merge paragraphs
                if t_idx == 5 and c_idx == 0 and r_idx > 0:
                    combined = " ".join([p.text.strip() for p in cell.paragraphs]).strip()
                    cell.text = combined
                
                # Determine alignment and boldness
                if r_idx == 0:
                    align = 'center'
                    bold = True
                else:
                    align = cfg['col_aligns'][c_idx] if c_idx < len(cfg['col_aligns']) else 'left'
                    bold = cfg['col_bold'][c_idx] if c_idx < len(cfg['col_bold']) else False
                
                # Format all paragraphs in the cell
                font_sz = cfg.get('font_size', 12)
                for p in cell.paragraphs:
                    format_cell_paragraph(p, align=align, bold=bold, font_size_pt=font_sz)
                    
    doc.save(file_path)
    print(f"Saved successfully: {file_path}")

if __name__ == '__main__':
    fix_all_tables_in_doc('baocaobaithi.docx')
